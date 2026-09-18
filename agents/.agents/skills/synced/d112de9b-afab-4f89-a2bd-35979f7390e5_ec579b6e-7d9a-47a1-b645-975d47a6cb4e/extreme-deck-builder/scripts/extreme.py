"""
extreme.py — helpers for building Extreme Networks branded decks.

The recipe:
  1. new_deck() copies the layouts-only base template (61 branded layouts,
     all chrome — logo, header swoosh, footer, page number — lives in the
     layouts, NOT in these helpers).
  2. add_slide(prs, "<Layout Name>", title=..., subtitle=...) picks a layout
     and fills its title placeholder(s).
  3. Add content as explicitly positioned shapes/text using the helpers
     below. Do NOT fight content placeholders by index — position freely
     inside CONTENT area.
  4. Save and run visual QA (see SKILL.md).

Usage:
    from extreme import *
    prs = new_deck()
    s = add_slide(prs, "Cover-1", title="Claude Usage Report\nMay 2026",
                  subtitle="Executive Summary", author="IT / Data Intelligence")
    s = add_slide(prs, "Title Only", title="Executive Summary \u2014 May 2026")
    add_stat(s, 0.6, 1.6, 3.8, value="26.3B", label="Total Tokens",
             sub="\u25b2 52% vs April")
    prs.save("deck.pptx")
"""
import copy
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ---------------------------------------------------------------- constants

BASE = Path(__file__).resolve().parent.parent / "assets" / "extreme_base.pptx"

# Brand palette (theme colors of the template)
PURPLE      = RGBColor(0x5B, 0x05, 0x9C)  # accent1 — primary brand purple
VIOLET      = RGBColor(0x75, 0x19, 0xF9)  # accent4 — bright violet
PERIWINKLE  = RGBColor(0x7D, 0x76, 0xF2)  # accent2
SKY         = RGBColor(0x9C, 0xE0, 0xF1)  # accent3 — pale cyan (use on dark)
DEEP        = RGBColor(0x24, 0x06, 0x91)  # accent5 — deep indigo
BLUE        = RGBColor(0x0F, 0x82, 0xF3)  # accent6
INK         = RGBColor(0x20, 0x00, 0x4C)  # dk2 — near-black purple (headings on light)
BLACK       = RGBColor(0x00, 0x00, 0x00)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
GRAY        = RGBColor(0x59, 0x59, 0x59)  # body gray on light slides
LIGHT_GRAY  = RGBColor(0xBF, 0xBF, 0xBF)  # secondary text on dark slides
CARD_DARK   = RGBColor(0x16, 0x10, 0x2E)  # card fill on dark slides
CARD_LIGHT  = RGBColor(0xF4, 0xF1, 0xFA)  # card fill on light slides

FONT = "DM Sans"   # template theme font (falls back gracefully if absent)

# Slide geometry (inches). Slide is 13.333 x 7.5.
SLIDE_W, SLIDE_H = 13.333, 7.5
# Safe content area on "Title Only"-family layouts (below the branded header
# band, above the footer/copyright line):
CX, CY, CX2, CY2 = 0.49, 1.35, 12.84, 6.95
CW, CH = CX2 - CX, CY2 - CY

# Layout names most decks need (full catalog: references/layouts.md)
COVER        = "Cover-1"                 # dark purple gradient cover
AGENDA       = "Agenda Slide"
LIGHT        = "Title Only"              # light content slide, branded header
LIGHT_SUB    = "Title Only Subhead"      # light, title + subtitle
DARK         = "Title Only Dark"         # dark content slide, branded header
DARK_SUB     = "Title Only Dark Subhead"
TWO_COL      = "1_Two Column"            # light, title + 2 text columns
THREE_COL    = "1_Three Column"          # light, title + 3 text columns
QUOTE        = "1_Content_Quote"         # big stat / customer quote
BLANK_LIGHT  = "4_Light-No-title"        # light, no chrome text
END          = "End"                     # dark closing slide with logo orb

# ------------------------------------------------------------ deck theme
# A deck uses ONE background theme for all content slides — dark or light.
# Ask the user which they want before building, then call set_theme() once.
# Cover and End slides always use the official dark brand artwork (the
# template provides no light variants of these); the theme governs every
# slide in between.

_THEME = "light"

def set_theme(theme: str):
    """Call once, right after new_deck(): set_theme('dark') or ('light')."""
    global _THEME
    if theme not in ("light", "dark"):
        raise ValueError("theme must be 'light' or 'dark'")
    _THEME = theme

def is_dark() -> bool:
    return _THEME == "dark"

def content_layout(subhead=False) -> str:
    """Theme-correct Title Only layout name."""
    if is_dark():
        return DARK_SUB if subhead else DARK
    return LIGHT_SUB if subhead else LIGHT

def column_layout(cols=2, subhead=False) -> str:
    """Theme-correct Two/Three Column layout name."""
    base = "1_Two Column" if cols == 2 else "1_Three Column"
    if subhead:
        base += " Subhead"
    return base + (" Dark" if is_dark() else "")

# Theme-aware color accessors — use these instead of hardcoding:
def c_head():   return WHITE if is_dark() else INK         # headings/values
def c_body():   return WHITE if is_dark() else GRAY        # body text
def c_muted():  return LIGHT_GRAY if is_dark() else GRAY   # captions/subs
def c_label():  return SKY if is_dark() else PURPLE        # section labels
def c_card():   return CARD_DARK if is_dark() else CARD_LIGHT
def c_bar():    return VIOLET                                # works on both


# ---------------------------------------------------------------- core

def new_deck(base: str | Path = BASE) -> Presentation:
    return Presentation(str(base))


def _layout(prs: Presentation, name: str):
    for master in prs.slide_masters:
        for lay in master.slide_layouts:
            if lay.name == name:
                return lay
    raise KeyError(f"No layout named {name!r}. See references/layouts.md")


def add_slide(prs: Presentation, layout_name: str, title: str | None = None,
              subtitle: str | None = None, author: str | None = None,
              keep_placeholders: bool = False):
    """Add a slide on the named layout and fill its title placeholder(s).

    Title placeholders are identified by position (topmost = title, next =
    subtitle) so this works across the whole layout family without knowing
    placeholder idx values. Unfilled placeholders are removed unless
    keep_placeholders=True (keep them when you plan to write into the
    column placeholders of Two/Three Column layouts via fill_columns()).
    """
    slide = prs.slides.add_slide(_layout(prs, layout_name))
    phs = sorted(slide.placeholders, key=lambda p: (p.top or 0, p.left or 0))
    texts = [t for t in (title, subtitle, author) if t is not None]
    filled = []
    for ph in phs:
        if texts:
            ph.text_frame.text = texts.pop(0)
            filled.append(ph)
    if not keep_placeholders:
        for ph in phs:
            if ph not in filled:
                ph._element.getparent().remove(ph._element)
    return slide


def fill_columns(slide, columns: list[list[tuple]]):
    """Fill the OBJECT placeholders of Two/Three Column layouts.

    columns = one list per column; each entry is (text, level) or
    (text, level, dict_of_run_kwargs). Requires the slide to have been
    created with keep_placeholders=True.
    """
    cols = sorted((p for p in slide.placeholders if p.placeholder_format.idx >= 16),
                  key=lambda p: p.left)
    for ph, items in zip(cols, columns):
        tf = ph.text_frame
        tf.clear()
        for i, item in enumerate(items):
            text, level, kw = (item + ({},))[:3] if len(item) == 2 else item
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.level = level
            run = para.add_run()
            run.text = text
            _style(run, **kw)


# ---------------------------------------------------------------- text

def _style(run, size=None, bold=None, color=None, font=FONT, italic=None):
    f = run.font
    f.name = font
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.rgb = color


def add_text(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, space_after=4, wrap=True):
    """Free text box. `lines` is a list of paragraphs; each paragraph is a
    list of run dicts: {"text": ..., "size": 14, "bold": True, "color": INK}.
    A paragraph dict may also carry {"bullet": "\u2022 "} prefix-style or
    "space_before": pts.
    Shorthand: a plain string is accepted for a one-run paragraph.
    """
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, para in enumerate(lines):
        if isinstance(para, str):
            para = [{"text": para}]
        if isinstance(para, dict):
            para = [para]
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        if para and "space_before" in para[0]:
            p.space_before = Pt(para[0]["space_before"])
        for rd in para:
            run = p.add_run()
            run.text = rd.get("bullet", "") + rd["text"]
            _style(run, size=rd.get("size"), bold=rd.get("bold"),
                   color=rd.get("color"), italic=rd.get("italic"),
                   font=rd.get("font", FONT))
    return box


# ---------------------------------------------------------------- blocks

def add_card(slide, x, y, w, h, fill=None, line=None, radius=0.08):
    """Rounded-rect card. Returns the shape (add text on top with add_text)."""
    from pptx.enum.shapes import MSO_SHAPE
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = radius
    shp.fill.solid()
    shp.fill.fore_color.rgb = c_card() if fill is None else fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    shp.shadow.inherit = False
    return shp


def add_stat(slide, x, y, w, value="", label="", sub="", dark=None, h=1.7):
    """KPI stat card: big number, label, small delta/sub line.
    `dark` defaults to the deck theme set via set_theme()."""
    dark = is_dark() if dark is None else dark
    fill = CARD_DARK if dark else CARD_LIGHT
    vcol = WHITE if dark else INK
    lcol = SKY if dark else PURPLE
    scol = LIGHT_GRAY if dark else GRAY
    add_card(slide, x, y, w, h, fill=fill)
    add_text(slide, x + 0.25, y + 0.18, w - 0.5, h - 0.36, [
        [{"text": value, "size": 32, "bold": True, "color": vcol}],
        [{"text": label, "size": 13, "bold": True, "color": lcol}],
        [{"text": sub, "size": 10.5, "color": scol}],
    ], space_after=2)


def add_bar_row(slide, x, y, w, label, frac, value_text, dark=None,
                bar_color=VIOLET, label_w=1.7, h=0.32):
    """One horizontal bar with left label and right value — for simple
    in-slide comparisons without a chart object.
    `dark` defaults to the deck theme set via set_theme()."""
    from pptx.enum.shapes import MSO_SHAPE
    dark = is_dark() if dark is None else dark
    tcol = WHITE if dark else INK
    add_text(slide, x, y + 0.02, label_w, h, [
        [{"text": label, "size": 10.5, "color": tcol}]])
    track_x, track_w = x + label_w + 0.1, w - label_w - 1.0
    bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                 Inches(track_x), Inches(y),
                                 Inches(max(track_w * min(frac, 1.0), 0.06)),
                                 Inches(h * 0.72))
    bar.adjustments[0] = 0.5
    bar.fill.solid(); bar.fill.fore_color.rgb = bar_color
    bar.line.fill.background(); bar.shadow.inherit = False
    add_text(slide, track_x + track_w + 0.12, y + 0.02, 0.8, h, [
        [{"text": value_text, "size": 10.5, "bold": True, "color": tcol}]])


def add_picture(slide, path, x, y, w=None, h=None):
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y),
                                    Inches(w) if w else None,
                                    Inches(h) if h else None)
